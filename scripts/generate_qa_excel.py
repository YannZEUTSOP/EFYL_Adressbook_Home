"""Generate the bilingual QA test script Excel file for EFYL Address Book."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = "/app/memory/EFYL_QA_Test_Script_FR_EN.xlsx"


HEADER_FR = ["ID", "Module", "Étape", "Action utilisateur", "Résultat attendu", "Résultat obtenu", "Statut (OK/KO)", "Sévérité", "Commentaires"]
HEADER_EN = ["ID", "Module", "Step", "User action", "Expected result", "Actual result", "Status (PASS/FAIL)", "Severity", "Comments"]

SCENARIOS_FR = [
    # Onboarding
    ("ONB-001", "Onboarding", 1, "Lancer l'app pour la première fois", "Carrousel 3 slides s'affiche en français", "", "", "Bloquant", ""),
    ("ONB-002", "Onboarding", 2, "Cliquer sur le bouton EN en haut à gauche", "Toute l'interface bascule en anglais", "", "", "Majeur", ""),
    ("ONB-003", "Onboarding", 3, "Cliquer sur Suivant 3 fois jusqu'à 'Commencer'", "Redirige vers la page principale (Accueil ou Login)", "", "", "Bloquant", ""),
    # Auth
    ("AUTH-001", "Authentification", 1, "Saisir +237 600 000 000 et appuyer sur 'Envoyer le code'", "Réception du code démo 123456 affiché à l'écran", "", "", "Bloquant", "Mode démo uniquement"),
    ("AUTH-002", "Authentification", 2, "Saisir 123456 dans le champ OTP", "Connexion réussie, redirection vers l'Accueil", "", "", "Bloquant", ""),
    ("AUTH-003", "Authentification", 3, "Tenter de saisir un mauvais code (ex: 000000)", "Message d'erreur 'Code invalide' s'affiche", "", "", "Majeur", ""),
    ("AUTH-004", "Authentification", 4, "Appuyer 4 fois 'Envoyer le code' en moins de 10 min", "Au 4ème essai : message 'Trop de requêtes, attendez 10 min'", "", "", "Majeur", "Rate limit 3/10min"),
    ("AUTH-005", "Authentification", 5, "Sur l'écran login, appuyer 'Continuer en invité'", "Accès à l'app sans login, mode local-only", "", "", "Mineur", ""),
    # Home
    ("HOME-001", "Accueil", 1, "Observer l'Accueil après login (sans adresse créée)", "État vide avec image + texte 'Aucune adresse'", "", "", "Mineur", ""),
    ("HOME-002", "Accueil", 2, "Observer le badge 'Hors ligne' en activant le mode avion", "Bandeau jaune 'Hors ligne — données sauvegardées localement'", "", "", "Majeur", ""),
    ("HOME-003", "Accueil", 3, "Vérifier la présence du FAB orange 'Nouvelle Adresse'", "Bouton flottant visible en bas à droite", "", "", "Bloquant", ""),
    ("HOME-004", "Accueil", 4, "Vérifier l'apparition de la section 'Commerces de la communauté'", "Liste s'affiche si commerces existent ou hidden si réseau down", "", "", "Majeur", "Nouveau"),
    # New address wizard
    ("ADR-001", "Création adresse", 1, "Appuyer sur le FAB 'Nouvelle Adresse'", "Wizard étape 1 'Mesurer la porte' s'affiche", "", "", "Bloquant", ""),
    ("ADR-002", "Création adresse", 2, "Autoriser permission GPS quand demandé", "Mesure démarre, satellites visibles, précision en mètres affichée", "", "", "Bloquant", ""),
    ("ADR-003", "Création adresse", 3, "Attendre une précision ≤ 1 m (pin vert)", "Bouton 'Confirmer la Position GPS' devient actif", "", "", "Majeur", ""),
    ("ADR-004", "Création adresse", 4, "Si précision plafonne >5m, appuyer 'Sauvegarder quand même'", "Adresse créée avec flag 'low_accuracy' rouge", "", "", "Majeur", ""),
    ("ADR-005", "Création adresse", 5, "Remplir Quartier/Secteur/Rue/Numéro, valider", "Passage à l'étape 3 'Direction & Photo'", "", "", "Majeur", "Quartier auto-détecté"),
    ("ADR-006", "Création adresse", 6, "Pointer le tél vers la porte, appuyer 'Verrouiller la direction'", "Azimut verrouillé, dégrés affichés", "", "", "Majeur", ""),
    ("ADR-007", "Création adresse", 7, "Prendre une photo du bâtiment (optionnel)", "Aperçu photo, bouton 'Reprendre'", "", "", "Mineur", ""),
    ("ADR-008", "Création adresse", 8, "Étape 4 : déplacer point orange sur image satellite", "Coordonnées corrigées sauvegardées séparément", "", "", "Mineur", ""),
    ("ADR-009", "Création adresse", 9, "Étape 5 : vérifier aperçu et appuyer 'Enregistrer'", "Adresse enregistrée, retour à l'accueil avec adresse listée", "", "", "Bloquant", ""),
    # Survey bifurcation
    ("SUR-001", "Enquête", 1, "Sur la page détail d'une adresse, appuyer 'Démarrer enquête'", "Question 'Foyer ou Commerce ?' avec 2 cartes", "", "", "Bloquant", "NOUVEAU"),
    ("SUR-002", "Enquête", 2, "Choisir 'Foyer (maison d'habitation)'", "Wizard 19 étapes Foyer démarre (Intro → Personnes)", "", "", "Bloquant", ""),
    ("SUR-003", "Enquête", 3, "Compléter Foyer jusqu'au récapitulatif", "Récapitulatif privé Foyer s'affiche avec tag bleu primary", "", "", "Bloquant", ""),
    ("SUR-004", "Enquête", 4, "Enregistrer l'enquête Foyer", "Adresse marquée has_survey=true mais is_commerce=false", "", "", "Bloquant", ""),
    ("SUR-005", "Enquête", 5, "Sur une autre adresse, choisir 'Commerce'", "Wizard 3 étapes Commerce démarre (choix → infos → résumé)", "", "", "Bloquant", "NOUVEAU"),
    ("SUR-006", "Enquête", 6, "Remplir Nom commerce, Propriétaire, Téléphone, Produits, Horaires", "Toutes les fenêtres acceptent texte", "", "", "Majeur", ""),
    ("SUR-007", "Enquête", 7, "Enregistrer enquête Commerce", "Adresse marquée is_commerce=true, tag orange Commerce visible sur la carte d'adresse", "", "", "Bloquant", "NOUVEAU"),
    # Community map
    ("MAP-001", "Carte", 1, "Aller à l'onglet Carte", "Carte Leaflet centrée sur les adresses utilisateur", "", "", "Bloquant", ""),
    ("MAP-002", "Carte", 2, "Activer le toggle Commerces communauté", "Pins oranges 🏪 apparaissent pour les commerces partagés par TOUS les utilisateurs", "", "", "Majeur", "NOUVEAU"),
    ("MAP-003", "Carte", 3, "Désactiver le toggle Commerces communauté", "Pins oranges des commerces externes disparaissent", "", "", "Majeur", "NOUVEAU"),
    ("MAP-004", "Carte", 4, "Basculer Plan / Satellite", "Tuiles satellite (Esri) remplacent OSM", "", "", "Mineur", ""),
    ("MAP-005", "Carte", 5, "Appuyer sur un pin", "Redirige vers la page détail de l'adresse", "", "", "Majeur", ""),
    # QR & Sharing
    ("QR-001", "QR Code", 1, "Sur le détail d'une adresse, appuyer 'QR Code'", "QR code généré et affiché", "", "", "Majeur", ""),
    ("QR-002", "QR Code", 2, "Appuyer 'Partager via WhatsApp'", "WhatsApp s'ouvre avec lien et code adresse pré-rempli", "", "", "Majeur", ""),
    ("QR-003", "Scan QR", 3, "Onglet Scanner, autoriser caméra, scanner QR EFYL", "Adresse trouvée, options 'Ajouter au carnet' et 'Itinéraire'", "", "", "Majeur", ""),
    # Verification
    ("VER-001", "Vérification", 1, "Se rendre près d'une adresse, appuyer 'Signaler une imprécision'", "Écran 'Vérifier l'adresse' avec position GPS live", "", "", "Majeur", ""),
    ("VER-002", "Vérification", 2, "Appuyer 'Soumettre' depuis la bonne position (< 10 m)", "Message 'Position vérifiée — merci !'", "", "", "Majeur", ""),
    ("VER-003", "Vérification", 3, "Appuyer 'Soumettre' depuis une fausse position (> 10 m)", "Message 'Position signalée — merci pour le retour'", "", "", "Majeur", ""),
    # Sync
    ("SYNC-001", "Synchronisation", 1, "Créer une adresse en mode avion (offline)", "Adresse sauvegardée localement, file de sync = 1", "", "", "Bloquant", ""),
    ("SYNC-002", "Synchronisation", 2, "Désactiver mode avion", "Sync automatique se déclenche, file = 0, dernière sync mise à jour", "", "", "Bloquant", ""),
    ("SYNC-003", "Synchronisation", 3, "Aller dans Profil > Synchroniser maintenant", "Sync manuelle réussie", "", "", "Majeur", ""),
    # Export
    ("EXP-001", "Export", 1, "Profil > Exporter mes données (CSV)", "Fichier CSV téléchargé avec toutes les colonnes adresse + enquête", "", "", "Majeur", ""),
    ("EXP-002", "Export", 2, "Profil > Exporter mes données (JSON)", "Fichier JSON téléchargé", "", "", "Mineur", ""),
    # Profile & language
    ("PRF-001", "Profil", 1, "Changer langue FR → EN dans Profil", "Toute l'app bascule en anglais immédiatement", "", "", "Majeur", ""),
    ("PRF-002", "Profil", 2, "Appuyer 'Supprimer mes données'", "Confirmation puis vidage AsyncStorage", "", "", "Majeur", ""),
    ("PRF-003", "Profil", 3, "Vérifier statistiques (X adresses créées, Y enquêtes)", "Chiffres exacts vs base de données", "", "", "Mineur", ""),
    # Edge cases
    ("EDG-001", "Cas limites", 1, "Créer une adresse hors du Cameroun (ex: Paris)", "Géocodage inverse fonctionne globalement (Nominatim)", "", "", "Majeur", ""),
    ("EDG-002", "Cas limites", 2, "Tenter de créer 2 adresses au même endroit", "Permis, mais adresses distinctes avec address_id différent", "", "", "Mineur", ""),
    ("EDG-003", "Cas limites", 3, "Supprimer une adresse depuis le détail", "Adresse retirée de la liste, retour à l'accueil", "", "", "Majeur", ""),
    ("EDG-004", "Cas limites", 4, "Ouvrir l'app après un crash brutal", "Données AsyncStorage intactes, file de sync préservée", "", "", "Majeur", ""),
]

SCENARIOS_EN = [
    ("ONB-001", "Onboarding", 1, "Launch app for the first time", "3-slide carousel displays in English", "", "", "Blocker", ""),
    ("ONB-002", "Onboarding", 2, "Tap FR button at top left", "Whole UI switches to French", "", "", "Major", ""),
    ("ONB-003", "Onboarding", 3, "Tap Next 3 times until 'Get Started'", "Redirects to main page (Home or Login)", "", "", "Blocker", ""),
    ("AUTH-001", "Authentication", 1, "Enter +237 600 000 000 and tap 'Send code'", "Demo code 123456 received and shown on screen", "", "", "Blocker", "Demo mode only"),
    ("AUTH-002", "Authentication", 2, "Enter 123456 in OTP field", "Login successful, redirect to Home", "", "", "Blocker", ""),
    ("AUTH-003", "Authentication", 3, "Try entering wrong code (e.g. 000000)", "Error 'Invalid code' displayed", "", "", "Major", ""),
    ("AUTH-004", "Authentication", 4, "Tap 'Send code' 4 times in less than 10 min", "On 4th attempt: 'Too many requests, wait 10 min'", "", "", "Major", "Rate limit 3/10min"),
    ("AUTH-005", "Authentication", 5, "On login screen tap 'Continue as guest'", "Access app without login, local-only mode", "", "", "Minor", ""),
    ("HOME-001", "Home", 1, "Observe Home after login (no address created)", "Empty state with image + 'No saved addresses'", "", "", "Minor", ""),
    ("HOME-002", "Home", 2, "Observe 'Offline' badge by enabling airplane mode", "Yellow banner 'Offline — data is safe locally'", "", "", "Major", ""),
    ("HOME-003", "Home", 3, "Verify presence of orange 'New Address' FAB", "Floating button visible bottom-right", "", "", "Blocker", ""),
    ("HOME-004", "Home", 4, "Verify 'Community commerces' section appears", "List shows if commerces exist, hidden if network down", "", "", "Major", "NEW"),
    ("ADR-001", "New Address", 1, "Tap 'New Address' FAB", "Wizard step 1 'Measure Doorstep' shown", "", "", "Blocker", ""),
    ("ADR-002", "New Address", 2, "Grant GPS permission when prompted", "Measurement starts, satellites visible, accuracy in meters", "", "", "Blocker", ""),
    ("ADR-003", "New Address", 3, "Wait for accuracy ≤ 1 m (green pin)", "Button 'Confirm GPS Position' becomes active", "", "", "Major", ""),
    ("ADR-004", "New Address", 4, "If accuracy plateaus >5m, tap 'Save Anyway'", "Address created with red 'low_accuracy' flag", "", "", "Major", ""),
    ("ADR-005", "New Address", 5, "Fill District/Sector/Street/Number, validate", "Move to step 3 'Direction & Photo'", "", "", "Major", "District auto-detected"),
    ("ADR-006", "New Address", 6, "Point phone at door, tap 'Lock direction'", "Azimuth locked, degrees displayed", "", "", "Major", ""),
    ("ADR-007", "New Address", 7, "Take a building photo (optional)", "Photo preview, 'Retake' button", "", "", "Minor", ""),
    ("ADR-008", "New Address", 8, "Step 4: drag orange pin on satellite image", "Corrected coords saved separately", "", "", "Minor", ""),
    ("ADR-009", "New Address", 9, "Step 5: review preview and tap 'Save'", "Address saved, back to home with address listed", "", "", "Blocker", ""),
    ("SUR-001", "Survey", 1, "On address detail, tap 'Start survey'", "Question 'Household or Commerce?' with 2 cards", "", "", "Blocker", "NEW"),
    ("SUR-002", "Survey", 2, "Choose 'Household (home)'", "19-step Household wizard starts (Intro → Persons)", "", "", "Blocker", ""),
    ("SUR-003", "Survey", 3, "Complete Household to summary", "Private Household summary shown with primary blue tag", "", "", "Blocker", ""),
    ("SUR-004", "Survey", 4, "Save Household survey", "Address marked has_survey=true but is_commerce=false", "", "", "Blocker", ""),
    ("SUR-005", "Survey", 5, "On another address, choose 'Commerce'", "3-step Commerce wizard (choice → info → summary)", "", "", "Blocker", "NEW"),
    ("SUR-006", "Survey", 6, "Fill Business name, Owner, Phone, Products, Hours", "All fields accept text", "", "", "Major", ""),
    ("SUR-007", "Survey", 7, "Save Commerce survey", "Address marked is_commerce=true, orange Commerce tag visible", "", "", "Blocker", "NEW"),
    ("MAP-001", "Map", 1, "Go to Map tab", "Leaflet map centered on user's addresses", "", "", "Blocker", ""),
    ("MAP-002", "Map", 2, "Activate Community commerces toggle", "Orange 🏪 pins appear for commerces shared by ALL users", "", "", "Major", "NEW"),
    ("MAP-003", "Map", 3, "Deactivate Community commerces toggle", "External commerce pins disappear", "", "", "Major", "NEW"),
    ("MAP-004", "Map", 4, "Switch Streets / Satellite", "Esri satellite tiles replace OSM", "", "", "Minor", ""),
    ("MAP-005", "Map", 5, "Tap a pin", "Redirects to the address detail page", "", "", "Major", ""),
    ("QR-001", "QR Code", 1, "On address detail, tap 'QR Code'", "QR code generated and shown", "", "", "Major", ""),
    ("QR-002", "QR Code", 2, "Tap 'Share via WhatsApp'", "WhatsApp opens with link and address code prefilled", "", "", "Major", ""),
    ("QR-003", "Scan QR", 3, "Scan tab, grant camera, scan EFYL QR", "Address found, 'Add to book' & 'Directions' options", "", "", "Major", ""),
    ("VER-001", "Verification", 1, "Stand near an address, tap 'Report inaccuracy'", "'Verify Address' screen with live GPS position", "", "", "Major", ""),
    ("VER-002", "Verification", 2, "Tap 'Submit' from correct position (< 10 m)", "Message 'Position verified — thank you!'", "", "", "Major", ""),
    ("VER-003", "Verification", 3, "Tap 'Submit' from wrong position (> 10 m)", "Message 'Inaccuracy reported — thanks for feedback'", "", "", "Major", ""),
    ("SYNC-001", "Sync", 1, "Create address in airplane mode (offline)", "Address saved locally, sync queue = 1", "", "", "Blocker", ""),
    ("SYNC-002", "Sync", 2, "Disable airplane mode", "Auto-sync triggers, queue = 0, last sync updated", "", "", "Blocker", ""),
    ("SYNC-003", "Sync", 3, "Profile > Sync Now", "Manual sync successful", "", "", "Major", ""),
    ("EXP-001", "Export", 1, "Profile > Export my data (CSV)", "CSV file downloaded with all address + survey columns", "", "", "Major", ""),
    ("EXP-002", "Export", 2, "Profile > Export my data (JSON)", "JSON file downloaded", "", "", "Minor", ""),
    ("PRF-001", "Profile", 1, "Switch language FR → EN in Profile", "Whole app switches to English instantly", "", "", "Major", ""),
    ("PRF-002", "Profile", 2, "Tap 'Delete my data'", "Confirmation then AsyncStorage cleared", "", "", "Major", ""),
    ("PRF-003", "Profile", 3, "Check stats (X addresses created, Y surveys)", "Exact numbers vs database", "", "", "Minor", ""),
    ("EDG-001", "Edge cases", 1, "Create an address outside Cameroon (e.g. Paris)", "Reverse geocoding works globally (Nominatim)", "", "", "Major", ""),
    ("EDG-002", "Edge cases", 2, "Try creating 2 addresses at same spot", "Allowed but with distinct address_id", "", "", "Minor", ""),
    ("EDG-003", "Edge cases", 3, "Delete an address from detail screen", "Address removed from list, back to home", "", "", "Major", ""),
    ("EDG-004", "Edge cases", 4, "Reopen app after a hard crash", "AsyncStorage intact, sync queue preserved", "", "", "Major", ""),
]

INFO_FR = [
    ["EFYL Address Book — Script de tests QA (FR)", ""],
    ["Version", "1.0"],
    ["Date", "Juin 2025"],
    ["Build cible", "Expo SDK 53+ / FastAPI 0.110+"],
    ["Identifiants de test", "+237 600 000 000  |  OTP : 123456"],
    ["Backend de test", "https://<deploy-url>/api/"],
    ["", ""],
    ["Instructions", ""],
    ["Statut", "OK = test réussi, KO = échec, BLOQUÉ = pas testable"],
    ["Sévérité", "Bloquant / Majeur / Mineur"],
    ["", ""],
    ["Légende NOUVEAU", "Scénarios créés dans le sprint S4 (bifurcation Foyer/Commerce + carte communautaire)"],
]
INFO_EN = [
    ["EFYL Address Book — QA test script (EN)", ""],
    ["Version", "1.0"],
    ["Date", "June 2025"],
    ["Target build", "Expo SDK 53+ / FastAPI 0.110+"],
    ["Test credentials", "+237 600 000 000  |  OTP: 123456"],
    ["Test backend", "https://<deploy-url>/api/"],
    ["", ""],
    ["Instructions", ""],
    ["Status", "PASS = test passed, FAIL = failed, BLOCKED = not testable"],
    ["Severity", "Blocker / Major / Minor"],
    ["", ""],
    ["NEW tag", "Scenarios created in sprint S4 (Household/Commerce bifurcation + community map)"],
]


def style_sheet(ws, header):
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="0D5C3A", end_color="0D5C3A", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="C0C0C0"),
        right=Side(style="thin", color="C0C0C0"),
        top=Side(style="thin", color="C0C0C0"),
        bottom=Side(style="thin", color="C0C0C0"),
    )
    for col_idx, _ in enumerate(header, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 32
    # Default widths
    widths = [12, 18, 8, 32, 36, 22, 14, 12, 24]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def fill_rows(ws, rows):
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    sev_colors = {"Bloquant": "FCA5A5", "Blocker": "FCA5A5", "Majeur": "FED7AA", "Major": "FED7AA", "Mineur": "BBF7D0", "Minor": "BBF7D0"}
    alt_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
            if r % 2 == 0:
                cell.fill = alt_fill
            if c == 8 and val in sev_colors:
                cell.fill = PatternFill(start_color=sev_colors[val], end_color=sev_colors[val], fill_type="solid")
                cell.font = Font(bold=True)
        ws.row_dimensions[r].height = 38


def fill_info(ws, lines):
    title_font = Font(bold=True, size=16, color="0D5C3A")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 80
    ws.cell(row=1, column=1, value=lines[0][0]).font = title_font
    for i, (k, v) in enumerate(lines[1:], 3):
        c1 = ws.cell(row=i, column=1, value=k)
        c1.font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)


wb = Workbook()
ws_info = wb.active
ws_info.title = "Info"
fill_info(ws_info, INFO_FR + [["", ""]] + INFO_EN)

ws_fr = wb.create_sheet("Tests FR")
ws_fr.append(HEADER_FR)
style_sheet(ws_fr, HEADER_FR)
fill_rows(ws_fr, SCENARIOS_FR)
ws_fr.freeze_panes = "A2"

ws_en = wb.create_sheet("Tests EN")
ws_en.append(HEADER_EN)
style_sheet(ws_en, HEADER_EN)
fill_rows(ws_en, SCENARIOS_EN)
ws_en.freeze_panes = "A2"

wb.save(OUTPUT)
print(f"Saved: {OUTPUT}")
print(f"FR scenarios: {len(SCENARIOS_FR)}")
print(f"EN scenarios: {len(SCENARIOS_EN)}")
